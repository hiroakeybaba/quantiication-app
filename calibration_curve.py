import streamlit as st
import pandas as pd
import numpy as np
import re
import zipfile
from io import BytesIO
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches

st.title("検量線作成ツール")

uploaded_file = st.file_uploader("Excelファイルをアップロード", type=["xlsx"])

if uploaded_file is not None:

    df = pd.read_excel(uploaded_file)
    st.success("読込成功")

    df.columns = df.columns.str.strip()

    # =========================
    # 濃度取得
    # =========================
    if "濃度" in df.columns:
        df["concentration"] = pd.to_numeric(df["濃度"], errors="coerce")
    else:
        file_col = [c for c in df.columns if "ファイル" in c][0]

        def extract_conc(x):
            parts = str(x).split("_")
            for p in parts:
                if "pg" in p:
                    try:
                        return float(p.split("pg")[0])
                    except:
                        return np.nan
            return np.nan

        df["concentration"] = df[file_col].apply(extract_conc)

    df = df[df["concentration"] > 0]

    # =========================
    # 数値変換
    # =========================
    for col in df.columns:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    # =========================
    # IS列
    # =========================
    IS_candidates = [c for c in df.columns if "d4" in c]
    if len(IS_candidates) == 0:
        st.error("IS列が見つかりません")
        st.stop()

    IS_col = IS_candidates[0]
    st.write("IS列:", IS_col)

    # =========================
    # 重み
    # =========================
    df["weight"] = 1 / df["concentration"]

    target_cols = [
        c for c in df.columns
        if "面積" in c and c != IS_col
    ]

    results = []
    plot_buffers = []

    # =========================
    # 検量線
    # =========================
    for col in target_cols:

        tmp = df[["concentration", col, IS_col, "weight"]].copy()
        tmp = tmp[tmp[IS_col] > 0]

        tmp["response"] = tmp[col] / tmp[IS_col]
        tmp = tmp.dropna()

        if len(tmp) < 2:
            continue

        x = tmp["concentration"]
        y = tmp["response"]
        w = tmp["weight"]

        # =========================
        # 回帰
        # =========================
        coef = np.polyfit(x, y, 1, w=np.sqrt(w))
        slope, intercept = coef

        y_pred = slope * x + intercept

        # =========================
        # 加重R2
        # =========================
        y_mean = np.average(y, weights=w)
        r2 = 1 - np.sum(w * (y - y_pred)**2) / np.sum(w * (y - y_mean)**2)

        # =========================
        # LLOQ / ULOQ（濃度）
        # =========================
        LLOQ = x.min()
        ULOQ = x.max()

        # =========================
        # ✅ 実測Response（最重要）
        # =========================

        # LLOQ points
        LLOQ_response = tmp[tmp["concentration"] == LLOQ]["response"].mean()

        # ULOQ points
        ULOQ_response = tmp[tmp["concentration"] == ULOQ]["response"].mean()

        # =========================
        # 名前整形
        # =========================
        name_clean = col.replace(" : 面積", "").strip()
        safe_name = re.sub(r"[^A-Za-z0-9_-]", "_", name_clean)

        # =========================
        # グラフ
        # =========================
        fig, ax = plt.subplots()

        ax.scatter(x, y, label="Data")
        ax.plot(x, y_pred, color="red", label="Fit")

        # LOQライン
        ax.axhline(LLOQ_response, linestyle="--", color="green", label="LLOQ")
        ax.axhline(ULOQ_response, linestyle="--", color="orange", label="ULOQ")

        ax.set_title(name_clean)
        ax.set_xlabel("Concentration")
        ax.set_ylabel("Response (Area / IS)")

        ax.text(max(x)*0.5, max(y)*0.9, f"R2={r2:.4f}")
        ax.text(max(x)*0.5, max(y)*0.8, f"y={slope:.4f}x+{intercept:.4f}")

        ax.legend()

        buf = BytesIO()
        fig.savefig(buf, format="png")
        plt.close(fig)

        buf.seek(0)
        plot_buffers.append((safe_name + ".png", buf))

        # =========================
        # 保存
        # =========================
        results.append({
            "Compound": name_clean,
            "N": len(tmp),
            "Slope": slope,
            "Intercept": intercept,
            "R_squared": r2,
            "LLOQ": LLOQ,
            "ULOQ": ULOQ,
            "LLOQ_response": LLOQ_response,
            "ULOQ_response": ULOQ_response
        })

    result_df = pd.DataFrame(results)
    st.dataframe(result_df)

    # =========================
    # Excel
    # =========================
    excel_buffer = BytesIO()
    result_df.to_excel(excel_buffer, index=False)

    # =========================
    # PPT
    # =========================
    prs = Presentation()
    blank = prs.slide_layouts[6]

    idx = 0
    while idx < len(plot_buffers):

        slide = prs.slides.add_slide(blank)

        pos = [(0.5,0.5),(5.5,0.5),(0.5,4),(5.5,4)]

        for i in range(4):
            if idx >= len(plot_buffers):
                break

            name, img = plot_buffers[idx]

            slide.shapes.add_picture(
                img,
                Inches(pos[i][0]),
                Inches(pos[i][1]),
                width=Inches(4)
            )

            idx += 1

    ppt_buffer = BytesIO()
    prs.save(ppt_buffer)

    # =========================
    # ZIP
    # =========================
    zip_buffer = BytesIO()

    with zipfile.ZipFile(zip_buffer, "w") as z:
        z.writestr("results.xlsx", excel_buffer.getvalue())
        z.writestr("plots.pptx", ppt_buffer.getvalue())

        for name, buf in plot_buffers:
            z.writestr("plots/" + name, buf.getvalue())

    st.download_button(
        "📦 ZIPダウンロード",
        zip_buffer.getvalue(),
        "calibration_results.zip"
    )
